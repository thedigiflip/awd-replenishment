"""
每月滯銷庫存報表（Excel，5 個分頁）

  摘要 / Replenishment US / Replenishment UK / Replenishment DE / Excess Inventory

- 明細分頁的欄位與數值 = dashboard「⬇ 匯出 CSV」（同一份欄位定義：core/report_columns.py；
  同一個資料來源：ReplenishmentService / ExcessInventoryService）
  差別只有：Excess Inventory 分頁列出全部 ASIN（依滯銷金額排序），dashboard 匯出的是目前篩選結果
- 檔案存檔：<DuckDB 資料夾>/reports/<YYYY-MM>/MAGEASY_Excess_Report_<YYYY-MM>.xlsx（雲端 = /data/reports）
- n8n 每月 1 號 09:00（台北）呼叫 POST /reports/monthly → 下載檔案 → Gmail 寄出
"""

import asyncio
import os
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

import structlog
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from core.config import settings
from core.database import new_conn
from core.report_columns import EXCESS_COLUMNS, replenishment_headers
from services.excess_inventory_service import ExcessInventoryService
from services.replenishment_service import ReplenishmentService

log = structlog.get_logger()

TZ_TAIPEI = timezone(timedelta(hours=8))   # 台灣無夏令時間
STALE_DAYS = 3                              # 超過幾天沒更新就在報表 / 信件加警告

SITES = [("US", "ATVPDKIKX0DER"), ("UK", "A1F83G8C2ARO7P"), ("DE", "A1PA6795UKMFR9")]

# 金額欄位（Excel 顯示為 $#,##0.00）
MONEY_KEYS = {"excess_value", "cog"}

_HDR_FILL = PatternFill("solid", fgColor="1F2D3D")
_HDR_FONT = Font(bold=True, color="FFFFFF")
_TITLE_FONT = Font(bold=True, size=14)
_SEC_FONT = Font(bold=True, size=11, color="1F2D3D")
_WARN_FONT = Font(bold=True, color="C62828")
_SEC_FILL = PatternFill("solid", fgColor="ECEFF1")


def reports_dir() -> Path:
    return Path(os.path.dirname(settings.DUCKDB_PATH)) / "reports"


def default_period(today: date | None = None) -> str:
    """報表月份 = 上個月（1 號產生的是上個月的月報）"""
    today = today or datetime.now(TZ_TAIPEI).date()
    first = today.replace(day=1)
    return (first - timedelta(days=1)).strftime("%Y-%m")


def report_filename(period: str) -> str:
    return f"MAGEASY_Excess_Report_{period}.xlsx"


def report_path(period: str) -> Path:
    return reports_dir() / period / report_filename(period)


def _cell_value(v):
    """與 dashboard CSV 相同的值：None → 空白；True → 'Y'；False → 空白"""
    if v is None:
        return None
    if v is True:
        return "Y"
    if v is False:
        return None
    return v


def _usd(v) -> str:
    return f"${(v or 0):,.0f}"


class MonthlyReportService:

    def __init__(self):
        self._rep = ReplenishmentService()
        self._ex = ExcessInventoryService()

    async def generate(self, period: str | None = None) -> dict:
        period = period or default_period()
        now = datetime.now(TZ_TAIPEI)

        # 資料：與 dashboard 完全相同的服務
        controls = await self._rep.get_controls()
        cap_months = controls["total_cap_months"]["value"]
        rep = {}
        for code, mp in SITES:
            rep[code] = await self._rep.get_table(mp)
        ex = await self._ex.get_table()
        extra = await asyncio.to_thread(self._load_extra_sync, period, now.date())

        warnings = [f"{f['source']} 已 {f['age_days']} 天未更新（最後：{f['last']}）"
                    for f in extra["freshness"] if f["stale"]]

        summary = self._build_summary(period, rep, ex, extra, cap_months)

        path = report_path(period)
        path.parent.mkdir(parents=True, exist_ok=True)
        await asyncio.to_thread(self._write_xlsx, path, period, now, rep, ex, extra,
                                summary, warnings, cap_months)

        subject = (("⚠️ " if warnings else "")
                   + f"[MAGEASY] 滯銷庫存月報 {period}｜全通路滯銷 {_usd(summary['excess_value'])}")
        result = {
            "period": period,
            "generated_at": now.strftime("%Y-%m-%d %H:%M"),
            "filename": path.name,
            "path": str(path),
            "download_url": f"/reports/monthly/{period}/file",
            "warnings": warnings,
            "summary": summary,
            "email_subject": subject,
            "email_html": self._email_html(period, now, summary, warnings),
        }
        log.info("monthly_report.generated", period=period, path=str(path),
                 excess_value=summary["excess_value"], warnings=len(warnings))
        return result

    # ── 額外資料：新鮮度、上月月底快照 ─────────────────────────────────────
    @staticmethod
    def _load_extra_sync(period: str, today: date) -> dict:
        conn = new_conn()
        try:
            def one(sql, p=None):
                v = conn.execute(sql, p or []).fetchone()[0]
                if v is None:
                    return None
                return v.date() if isinstance(v, datetime) else v

            sources = [("FBA US", one("SELECT MAX(snapshot_date) FROM inventory WHERE marketplace_id='ATVPDKIKX0DER'")),
                       ("FBA UK", one("SELECT MAX(snapshot_date) FROM inventory WHERE marketplace_id='A1F83G8C2ARO7P'")),
                       ("FBA DE", one("SELECT MAX(snapshot_date) FROM inventory WHERE marketplace_id='A1PA6795UKMFR9'")),
                       ("AWD", one("SELECT MAX(synced_at) FROM awd_inventory")),
                       ("美國倉 / 佳樂倉（深圳倉 template）", one("SELECT MAX(synced_at) FROM sz_warehouse")),
                       ("銷量 US", one("SELECT MAX(report_date) FROM sales_summary WHERE marketplace_id='ATVPDKIKX0DER'")),
                       ("銷量 UK", one("SELECT MAX(report_date) FROM sales_summary WHERE marketplace_id='A1F83G8C2ARO7P'")),
                       ("銷量 DE", one("SELECT MAX(report_date) FROM sales_summary WHERE marketplace_id='A1PA6795UKMFR9'")),
                       ("MAGEASY Anchor", one("SELECT MAX(updated_at) FROM product_catalog"))]
            freshness = []
            for name, d in sources:
                age = (today - d).days if d else None
                # Anchor 不是每日資料，不列入過期警告
                stale = (d is None or age > STALE_DAYS) and name != "MAGEASY Anchor"
                freshness.append({"source": name, "last": str(d) if d else "無資料",
                                  "age_days": age if age is not None else "—", "stale": stale})

            # 上個月（相對報表月份）的月底快照，做月增減比較
            y, m = map(int, period.split("-"))
            prev_last = date(y, m, 1) - timedelta(days=1)
            prev_period = prev_last.strftime("%Y-%m")
            prev = conn.execute("""
                SELECT snapshot_date, SUM(excess_value), SUM(excess_qty)
                FROM fact_excess_inventory_daily
                WHERE snapshot_date = (SELECT MAX(snapshot_date) FROM fact_excess_inventory_daily
                                       WHERE strftime(snapshot_date, '%Y-%m') = ?)
                GROUP BY snapshot_date
            """, [prev_period]).fetchone()
            prev_snap = ({"period": prev_period, "as_of": str(prev[0]),
                          "excess_value": float(prev[1] or 0), "excess_qty": int(prev[2] or 0)}
                         if prev else {"period": prev_period, "as_of": None})
            return {"freshness": freshness, "prev": prev_snap}
        finally:
            conn.close()

    # ── 摘要數字（全部由明細資料計算，確保與分頁一致）──────────────────────
    @staticmethod
    def _build_summary(period, rep, ex, extra, cap_months) -> dict:
        s = ex["summary"]
        bm = (ex["params"].get("excess_benchmark_usd") or {}).get("value") or 0
        prev = extra["prev"]
        mom = None
        if prev.get("as_of"):
            mom = {"prev_as_of": prev["as_of"], "prev_value": prev["excess_value"],
                   "diff": s["excess_value"] - prev["excess_value"],
                   "pct": ((s["excess_value"] - prev["excess_value"]) / prev["excess_value"] * 100)
                          if prev["excess_value"] else None}
        locs = []
        for key, label in [("us_wh", "美國倉"), ("awd", "AWD"), ("fba_us", "FBA US"),
                           ("fba_uk", "FBA UK"), ("fba_de", "FBA DE")]:
            locs.append({"label": label, "value": round(s.get("value_" + key, 0) or 0, 2),
                         "qty": int(s.get("qty_" + key, 0) or 0)})
        locs.sort(key=lambda x: -x["value"])

        coll = {}
        for r in ex["data"]:
            if (r.get("excess_value") or 0) > 0:
                c = r.get("collection") or "(未分類)"
                a = coll.setdefault(c, {"collection": c, "value": 0.0, "qty": 0, "asins": 0})
                a["value"] += r["excess_value"]; a["qty"] += r["excess_qty"]; a["asins"] += 1
        for a in coll.values():
            a["value"] = round(a["value"], 2)
        top_coll = sorted(coll.values(), key=lambda x: -x["value"])[:10]

        sites = []
        for code, _ in SITES:
            rows = rep[code]
            sites.append({
                "site": code,
                "asins": len(rows),
                "amzn_excess_qty": sum(max(0, int(r.get("excess_qty") or 0)) for r in rows),
                "amzn_excess_value": round(sum(max(0.0, float(r.get("excess_value") or 0)) for r in rows), 2),
                "reorder": sum(1 for r in rows if "需返單" in (r.get("reorder_alert") or "")),
                "air": sum(1 for r in rows if r.get("air_alert")),
                "over_cap": sum(1 for r in rows if r.get("cap_alert") == "🔴 超過上限"),
            })
        return {
            "period": period,
            "excess_qty": s["excess_qty"], "excess_value": s["excess_value"],
            "excess_asins": s["excess_asins"], "stock_total": s["stock_total"],
            "still_replenishing": s["still_replenishing"], "potential_asins": s["potential_asins"],
            "benchmark": bm, "benchmark_pct": (s["excess_value"] / bm * 100) if bm else None,
            "annual_months": s.get("annual_months") or 12,
            "mom": mom, "locations": locs, "top_collections": top_coll, "sites": sites,
            "cap_months": cap_months,
        }

    # ── Excel ────────────────────────────────────────────────────────────────
    def _write_xlsx(self, path, period, now, rep, ex, extra, summary, warnings, cap_months):
        wb = Workbook()
        ws = wb.active
        ws.title = "摘要"
        self._write_summary_sheet(ws, period, now, extra, summary, warnings)

        rep_cols = replenishment_headers(cap_months)
        for code, _ in SITES:
            self._write_table(wb.create_sheet(f"Replenishment {code}"), rep_cols, rep[code], freeze="F2")
        self._write_table(wb.create_sheet("Excess Inventory"), EXCESS_COLUMNS, ex["data"], freeze="D2")

        tmp = path.with_suffix(".tmp.xlsx")
        wb.save(tmp)
        os.replace(tmp, path)     # 寫完才換名，避免 n8n 下載到寫一半的檔案

    @staticmethod
    def _write_table(ws, cols, rows, freeze):
        ws.append([h for _, h in cols])
        for cell in ws[1]:
            cell.fill = _HDR_FILL; cell.font = _HDR_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.row_dimensions[1].height = 32
        for r in rows:
            ws.append([_cell_value(r.get(k)) for k, _ in cols])
        for i, (k, h) in enumerate(cols, start=1):
            letter = get_column_letter(i)
            width = 40 if k == "product_name" else max(10, min(24, len(h) * 1.6 + 2))
            ws.column_dimensions[letter].width = width
            if k in MONEY_KEYS:
                for c in ws[letter][1:]:
                    c.number_format = "$#,##0.00"
        ws.freeze_panes = freeze
        ws.auto_filter.ref = ws.dimensions

    @staticmethod
    def _write_summary_sheet(ws, period, now, extra, s, warnings):
        ws.column_dimensions["A"].width = 34
        for col in "BCDEF":
            ws.column_dimensions[col].width = 18
        row = 1

        def put(values, font=None, fill=None, fmt=None):
            nonlocal row
            for i, v in enumerate(values, start=1):
                c = ws.cell(row=row, column=i, value=v)
                if font: c.font = font
                if fill: c.fill = fill
                if fmt and i > 1 and isinstance(v, (int, float)): c.number_format = fmt
            row += 1

        def section(title):
            nonlocal row
            row += 1
            put([title], font=_SEC_FONT, fill=_SEC_FILL)

        put([f"MAGEASY 滯銷庫存月報 — {period}"], font=_TITLE_FONT)
        put([f"產生時間：{now.strftime('%Y-%m-%d %H:%M')}（台北）｜資料為產生當下 dashboard 的數字"])
        put(["分頁：Replenishment US / UK / DE（= dashboard 匯出 CSV）、Excess Inventory（全部 ASIN，依滯銷金額排序）"])

        if warnings:
            section("⚠️ 資料新鮮度警告（數字可能不準，請先確認）")
            for w in warnings:
                put([w], font=_WARN_FONT)

        section("全通路滯銷（市管處定義）")
        put(["全通路滯銷金額 (USD)", s["excess_value"]], fmt="$#,##0")
        put(["全通路滯銷數量", s["excess_qty"]], fmt="#,##0")
        put(["滯銷 ASIN 數", s["excess_asins"]], fmt="#,##0")
        put(["庫存池合計（件）", s["stock_total"]], fmt="#,##0")
        if s["benchmark"]:
            put(["市管處 Benchmark (USD)", s["benchmark"]], fmt="$#,##0")
            put(["Benchmark 使用率", f"{s['benchmark_pct']:.0f}%"])
        else:
            put(["市管處 Benchmark (USD)", "尚未設定"])
        if s["mom"]:
            m = s["mom"]
            put([f"上月月底水位（快照 {m['prev_as_of']}）", m["prev_value"]], fmt="$#,##0")
            put(["月增減 (USD)", m["diff"]], fmt="$#,##0;-$#,##0")
            if m["pct"] is not None:
                put(["月增減 %", f"{m['pct']:+.1f}%"])
        else:
            put(["上月月底水位", "無快照資料（每日快照自 2026-10-26 開始）"])
        put(["⚠️ 滯銷仍在補貨（ASIN 數）", s["still_replenishing"]], fmt="#,##0")
        put(["🟡 潛在滯銷 9–12M（ASIN 數）", s["potential_asins"]], fmt="#,##0")

        section("滯銷在哪裡（各站先用自己的庫存，剩下的才是該位置的滯銷）")
        put(["位置", "滯銷金額 (USD)", "滯銷數量", "佔比"], font=Font(bold=True))
        tot = sum(l["value"] for l in s["locations"]) or 1
        for l in s["locations"]:
            put([l["label"], l["value"], l["qty"], f"{l['value'] / tot * 100:.0f}%"], fmt="#,##0")
            ws.cell(row=row - 1, column=2).number_format = "$#,##0"

        section("滯銷來源 — Collection 前 10 名（依金額）")
        put(["Collection", "滯銷金額 (USD)", "滯銷數量", "ASIN 數"], font=Font(bold=True))
        for cl in s["top_collections"]:
            put([cl["collection"], cl["value"], cl["qty"], cl["asins"]], fmt="#,##0")
            ws.cell(row=row - 1, column=2).number_format = "$#,##0"

        section(f"各站補貨重點（Replenishment；Cap = {s['cap_months']:g} 個月）")
        put(["站點", "AMZN 滯銷金額 (USD)", "AMZN 滯銷數量", "🔴 需返單", "🚨 空運警示", "🔴 超過 Cap"],
            font=Font(bold=True))
        for st in s["sites"]:
            put([st["site"], st["amzn_excess_value"], st["amzn_excess_qty"], st["reorder"], st["air"],
                 st["over_cap"]], fmt="#,##0")
            ws.cell(row=row - 1, column=2).number_format = "$#,##0"

        section("資料新鮮度")
        put(["資料來源", "最後更新", "距今天數", "狀態"], font=Font(bold=True))
        for f in extra["freshness"]:
            put([f["source"], f["last"], f["age_days"], "⚠️ 過期" if f["stale"] else "OK"],
                font=_WARN_FONT if f["stale"] else None)

        section("定義")
        months = s["annual_months"]
        for line in [
            f"全通路滯銷數量 = max(0, FBA Available(US+UK+DE) + AWD Available + 美國倉 − 全通路 30 天銷量 × {months:g})，逐 ASIN 計算後加總",
            "全通路滯銷金額 = 滯銷數量 × Anchor COG（USD）；全通路目前 = Amazon US + CA + MX + UK + DE",
            "AMZN 滯銷（Replenishment 分頁）= 單一站點：max(0, 該站 FBA Available + AWD Available − 該站 30 天銷量 × 12)",
            "US 站銷量含 CA + MX（北美遠程配送，由 US FBA 庫存出貨）",
        ]:
            put([line])

    # ── Email 內文 ───────────────────────────────────────────────────────────
    @staticmethod
    def _email_html(period, now, s, warnings) -> str:
        warn = ""
        if warnings:
            items = "".join(f"<li>{w}</li>" for w in warnings)
            warn = (f'<div style="background:#FDEEED;border-left:4px solid #E15759;padding:8px 12px;margin:10px 0">'
                    f'<b>⚠️ 資料新鮮度警告</b>（數字可能不準，請先確認）<ul style="margin:4px 0">{items}</ul></div>')
        bm = (f"{_usd(s['benchmark'])}（使用率 {s['benchmark_pct']:.0f}%）" if s["benchmark"] else "尚未設定")
        mom = "無上月快照"
        if s["mom"]:
            m = s["mom"]
            sign = "+" if m["diff"] >= 0 else "−"
            mom = f"{sign}${abs(m['diff']):,.0f}" + (f"（{m['pct']:+.1f}%）" if m["pct"] is not None else "")
        loc = "".join(f"<tr><td>{l['label']}</td><td style='text-align:right'>{_usd(l['value'])}</td>"
                      f"<td style='text-align:right'>{l['qty']:,}</td></tr>" for l in s["locations"])
        td = "padding:4px 10px;border-bottom:1px solid #eee"
        return f"""<div style="font-family:Arial,sans-serif;font-size:13px;color:#2C2C2C">
<h2 style="margin:0 0 6px">MAGEASY 滯銷庫存月報 — {period}</h2>
<div style="color:#777">產生時間 {now.strftime('%Y-%m-%d %H:%M')}（台北）</div>
{warn}
<table style="border-collapse:collapse;margin:10px 0">
<tr><td style="{td}">全通路滯銷金額</td><td style="{td}"><b>{_usd(s['excess_value'])}</b></td></tr>
<tr><td style="{td}">全通路滯銷數量</td><td style="{td}">{s['excess_qty']:,} 件（{s['excess_asins']:,} 個 ASIN）</td></tr>
<tr><td style="{td}">市管處 Benchmark</td><td style="{td}">{bm}</td></tr>
<tr><td style="{td}">與上月月底比較</td><td style="{td}">{mom}</td></tr>
</table>
<b>滯銷在哪裡</b>
<table style="border-collapse:collapse;margin:6px 0 12px">
<tr><th style="text-align:left;padding:4px 10px">位置</th><th style="padding:4px 10px">金額</th><th style="padding:4px 10px">數量</th></tr>
{loc}
</table>
<div>附件 Excel 共 5 個分頁：摘要、Replenishment US / UK / DE、Excess Inventory。<br>
明細欄位與 dashboard「匯出 CSV」相同，請協助核對數據。</div>
</div>"""
